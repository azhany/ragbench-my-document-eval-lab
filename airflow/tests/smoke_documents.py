"""Actual HTTP→Airflow smoke with synthetic fixtures; retains evidence for review.

Run inside the Airflow image with API_URL pointing at the Go API. Without a
provider key, --expect-missing-key verifies real extraction then the visible
embedding failure. Omit it to require a real-provider processed result.
"""
import argparse
import io
import json
import os
import time
import uuid
from urllib.error import HTTPError
from urllib.request import Request, urlopen


def request(method, path, body=None, content_type="application/json"):
    payload = json.dumps(body).encode() if body is not None and not isinstance(body, bytes) else body
    req = Request(os.getenv("API_URL", "http://backend:8080") + path, data=payload, method=method,
                  headers={"Content-Type": content_type})
    try:
        with urlopen(req, timeout=50) as response:
            return response.status, json.load(response) if response.status != 204 else None
    except HTTPError as exc:
        raise RuntimeError(f"HTTP {exc.code}: {exc.read().decode()}") from None


def upload(config_id, filename, payload):
    boundary = "rb-" + uuid.uuid4().hex
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"config_id\"\r\n\r\n{config_id}\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
            "Content-Type: application/octet-stream\r\n\r\n").encode() + payload + f"\r\n--{boundary}--\r\n".encode()
    status, doc = request("POST", "/api/v1/documents", body, f"multipart/form-data; boundary={boundary}")
    assert status == 202 and doc["status"] == "queued", doc
    print(json.dumps({"event": "uploaded", "document_id": doc["id"], "filename": filename, "run_id": doc["job"]["run_id"]}), flush=True)
    return doc


def wait_document(doc):
    deadline = time.monotonic() + 240
    while time.monotonic() < deadline:
        _, current = request("GET", "/api/v1/documents/" + doc["id"])
        if current["status"] == "processed" or (current["status"] == "failed" and current["job"]["error_code"]):
            print(json.dumps({"event": "finished", "document_id": current["id"], "status": current["status"],
                              "job": current["job"], "chunk_count": current["chunk_count"]}), flush=True)
            return current
        time.sleep(2)
    raise RuntimeError("Timed out waiting for document " + doc["id"])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--expect-missing-key", action="store_true")
    args = parser.parse_args()
    suffix = uuid.uuid4().hex[:8]
    _, config = request("POST", "/api/v1/rag-configs", {
        "name": "sprint2-smoke-" + suffix, "chunk_size": 80, "chunk_overlap": 10,
        "retrieval_mode": "vector", "top_k": 5, "prompt_version": "v1",
        "model_profile": "openai-gpt-4o-mini", "embedding_profile": "openai-text-embedding-3-small",
    })
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    text = f"Synthetic approval evidence {suffix}. A reviewer must approve a document before publication."
    docx = Document(); docx.add_heading("Approval policy", level=1); docx.add_paragraph(text)
    docx_bytes = io.BytesIO(); docx.save(docx_bytes)
    writer = PdfWriter(); page = writer.add_blank_page(width=600, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(f"BT /F1 10 Tf 20 200 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf_bytes = io.BytesIO(); writer.write(pdf_bytes)
    docs = [upload(config["id"], f"sprint2-{suffix}.{ext}", payload) for ext, payload in
            [("txt", text.encode()), ("docx", docx_bytes.getvalue()), ("pdf", pdf_bytes.getvalue())]]
    bad = upload(config["id"], f"sprint2-corrupt-{suffix}.pdf", b"corrupt synthetic PDF " + suffix.encode())
    for doc in docs:
        result = wait_document(doc)
        if args.expect_missing_key:
            assert result["job"]["error_code"] == "embedding_failed", result
            assert "not configured" in result["job"]["error_message"], result
            assert result["active_revision_id"] is None, result
            assert result["job"]["stage_metrics"]["chunk"]["status"] == "succeeded", result
        else:
            assert result["status"] == "processed" and result["chunk_count"] > 0, result
    assert wait_document(bad)["job"]["error_code"] == "corrupt_document"
    status, replacement = request("POST", f"/api/v1/documents/{docs[0]['id']}/reprocess", {"config_id": config["id"]})
    assert status == 202 and replacement["job"]["dag_id"] == "document_reindex", replacement
    if not args.expect_missing_key:
        previous_revision = replacement["active_revision_id"]
        assert previous_revision and previous_revision != replacement["latest_revision_id"], replacement
        replacement = wait_document(replacement)
        assert replacement["status"] == "processed", replacement
        assert replacement["active_revision_id"] == replacement["latest_revision_id"], replacement
        assert replacement["active_revision_id"] != previous_revision, replacement
        assert len([rev for rev in replacement["revisions"] if rev["status"] == "ready"]) == 2, replacement
        print(json.dumps({"event": "replacement_published", "document_id": replacement["id"],
                          "previous_revision": previous_revision, "active_revision": replacement["active_revision_id"]}), flush=True)
        # Queue a further revision to verify deletion while work is in flight.
        status, replacement = request("POST", f"/api/v1/documents/{docs[0]['id']}/reprocess", {"config_id": config["id"]})
        assert status == 202, replacement
    # Deleting a just-dispatched replacement exercises the live API cancellation path.
    assert request("DELETE", f"/api/v1/documents/{docs[0]['id']}")[0] == 204
    assert request("DELETE", f"/api/v1/documents/{docs[0]['id']}")[0] == 204
    _, listing = request("GET", "/api/v1/documents")
    assert docs[0]["id"] not in [doc["id"] for doc in listing["documents"]]
    print(json.dumps({"event": "smoke_passed", "real_provider": not args.expect_missing_key,
                      "config_id": config["id"], "documents": [doc["id"] for doc in docs],
                      "deleted_document": docs[0]["id"], "corrupt_document": bad["id"]}), flush=True)


if __name__ == "__main__":
    main()

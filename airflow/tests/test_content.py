import hashlib
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from ragbench.content import IngestionError, chunk, extract, normalize


class ContentTests(unittest.TestCase):
    def extract_file(self, path, mime):
        return extract(path, mime, hashlib.sha256(path.read_bytes()).hexdigest(), path.parent, 100000, 10000000)

    def test_exact_boundaries_overlap_and_retry_identity(self):
        revision = str(uuid.uuid4())
        normalized = normalize([{"text": "abcdefghij", "location": {"page": 2}}])
        result = chunk(normalized, revision, 4, 1)
        self.assertEqual([c["content"] for c in result], ["abcd", "defg", "ghij"])
        self.assertEqual([c["metadata"]["start"] for c in result], [0, 3, 6])
        self.assertEqual(result, chunk(normalized, revision, 4, 1))
        self.assertNotEqual(result[0]["id"], chunk(normalized, str(uuid.uuid4()), 4, 1)[0]["id"])
        self.assertEqual(result[0]["metadata"]["locations"][0]["page"], 2)
        with self.assertRaises(IngestionError):
            chunk(normalized, revision, 4, 4)

    def test_unicode_and_cross_page_locations(self):
        normalized = normalize([{"text": " e\u0301\x00  a ", "location": {"page": 1}},
                                {"text": " b\t c ", "location": {"page": 2}}])
        self.assertEqual(normalized["text"], "é a\nb c")
        result = chunk(normalized, str(uuid.uuid4()), 5, 0)
        self.assertEqual([loc["page"] for loc in result[0]["metadata"]["locations"]], [1, 2])

    def test_three_formats_and_classified_failures(self):
        from docx import Document
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            txt = root / "sample.txt"
            txt.write_text("Source approval policy.", encoding="utf-8")
            self.assertTrue(self.extract_file(txt, "text/plain"))
            docx = root / "sample.docx"
            doc = Document(); doc.add_heading("Policy", level=1); doc.add_paragraph("Approval needs a reviewer.")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
            doc.save(docx)
            sections = self.extract_file(docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.assertEqual(sections[-1]["text"], "Table evidence")
            self.assertEqual(sections[1]["location"]["section"], "Policy")
            pdf = root / "sample.pdf"
            writer = PdfWriter(); page = writer.add_blank_page(width=300, height=300)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
            stream = DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 20 200 Td (PDF source evidence.) Tj ET")
            page[NameObject("/Contents")] = writer._add_object(stream)
            writer.write(pdf)
            self.assertIn("PDF source", self.extract_file(pdf, "application/pdf")[0]["text"])
            writer.encrypt("secret"); writer.write(pdf)
            with self.assertRaises(IngestionError) as caught:
                self.extract_file(pdf, "application/pdf")
            self.assertEqual(caught.exception.code, "encrypted_document")
            writer = PdfWriter(); writer.add_blank_page(width=300, height=300); writer.write(pdf)
            with self.assertRaises(IngestionError) as caught:
                self.extract_file(pdf, "application/pdf")
            self.assertEqual(caught.exception.code, "image_only_or_empty")
            for payload, mime, code in [(b" \n", "text/plain", "empty_content"),
                                        (b"\xff", "text/plain", "corrupt_document"),
                                        (b"corrupt", "application/pdf", "corrupt_document"),
                                        (b"badzip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "corrupt_document")]:
                txt.write_bytes(payload)
                with self.assertRaises(IngestionError) as caught:
                    self.extract_file(txt, mime)
                self.assertEqual(caught.exception.code, code)

    def test_bounded_image_ocr_path_and_unreadable_fixture(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = root / "receipt.png"
            Image.new("RGB", (120, 80), "white").save(image)
            with patch("pytesseract.image_to_string", return_value="TOTAL MYR 10.00") as ocr:
                sections = self.extract_file(image, "image/png")
            self.assertIn("TOTAL MYR", sections[0]["text"])
            ocr.assert_called_once()

            corrupt = root / "corrupt.png"
            corrupt.write_bytes(b"not an image")
            with self.assertRaises(IngestionError) as caught:
                self.extract_file(corrupt, "image/png")
            self.assertEqual(caught.exception.code, "image_unreadable")


if __name__ == "__main__":
    unittest.main()
